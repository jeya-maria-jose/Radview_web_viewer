from __future__ import annotations

import os
import re
import uuid
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, List, Tuple

import docx
import nibabel as nib
import numpy as np
import pydicom
from flask import Flask, Response, jsonify, request, send_from_directory
from pydicom.errors import InvalidDicomError
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"

app = Flask(__name__, static_folder=str(STATIC_DIR), static_url_path="")


@dataclass
class Segmentation:
    mask_id: str
    source: str
    data: np.ndarray
    classes: List[dict]


@dataclass
class Volume:
    volume_id: str
    source: str
    kind: str
    data: np.ndarray
    spacing: Tuple[float, float, float]
    intensity_min: float
    intensity_max: float
    segmentations: Dict[str, Segmentation] = field(default_factory=dict)


VOLUMES: Dict[str, Volume] = {}


IMPORTANT_FINDING_RULES = [
    ("critical", "critical result"),
    ("critical", "urgent"),
    ("critical", "stat"),
    ("high", "acute"),
    ("high", "hemorrhage"),
    ("high", "haemorrhage"),
    ("high", "infarct"),
    ("high", "ischemia"),
    ("high", "ischaemia"),
    ("high", "embol"),
    ("high", "pneumothorax"),
    ("high", "midline shift"),
    ("high", "herniation"),
    ("high", "aneurysm"),
    ("high", "dissection"),
    ("high", "abscess"),
    ("high", "obstruction"),
    ("high", "perforation"),
    ("high", "fracture"),
    ("moderate", "mass"),
    ("moderate", "lesion"),
    ("moderate", "nodule"),
    ("moderate", "tumor"),
    ("moderate", "tumour"),
    ("moderate", "metast"),
    ("moderate", "malignan"),
    ("moderate", "edema"),
    ("moderate", "oedema"),
    ("moderate", "stenosis"),
    ("moderate", "enlarged"),
    ("moderate", "lymphadenopathy"),
]


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: List[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() in {"br", "p", "div", "tr", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def text(self) -> str:
        return "\n".join(part for part in self.parts if part)


def _resolve_user_path(raw_path: str) -> Path:
    if not raw_path or not raw_path.strip():
        raise ValueError("Path is required.")

    expanded = os.path.expandvars(os.path.expanduser(raw_path.strip()))
    path = Path(expanded).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Path does not exist: {path}")
    return path


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _strip_html(text: str) -> str:
    parser = _TextHTMLParser()
    parser.feed(text)
    return parser.text()


def _strip_rtf(text: str) -> str:
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+\n", "\n", text)


def _extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"Page {index}\n{text}")
    return "\n\n".join(pages)


def _extract_docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _collect_dicom_text(dataset: pydicom.Dataset) -> List[str]:
    parts = []
    for element in dataset:
        if element.VR == "SQ":
            for item in element.value:
                parts.extend(_collect_dicom_text(item))
            continue
        if element.keyword in {"TextValue", "LongText", "ShortText", "PersonName"}:
            value = str(element.value).strip()
            if value:
                parts.append(value)
        elif element.keyword == "ConceptNameCodeSequence":
            for item in element.value:
                meaning = str(getattr(item, "CodeMeaning", "")).strip()
                if meaning:
                    parts.append(meaning)
    return parts


def _extract_dicom_report_text(path: Path) -> str:
    dataset = pydicom.dcmread(str(path), stop_before_pixels=True, force=True)
    parts = []
    for keyword in ("StudyDescription", "SeriesDescription", "ProtocolName"):
        value = str(getattr(dataset, keyword, "")).strip()
        if value:
            parts.append(value)
    parts.extend(_collect_dicom_text(dataset))
    return "\n".join(dict.fromkeys(parts))


def _extract_report_text(path: Path) -> Tuple[str, str]:
    lower_name = path.name.lower()
    suffix = path.suffix.lower()

    if lower_name.endswith(".pdf"):
        return _extract_pdf_text(path), "pdf"
    if lower_name.endswith(".docx"):
        return _extract_docx_text(path), "docx"
    if suffix in {".html", ".htm"}:
        return _strip_html(_read_text_file(path)), "html"
    if suffix == ".rtf":
        return _strip_rtf(_read_text_file(path)), "rtf"
    if suffix in {".dcm", ".dicom"}:
        return _extract_dicom_report_text(path), "dicom-report"
    if suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".log", ".report"}:
        return _read_text_file(path), suffix.lstrip(".") or "text"

    try:
        return _extract_dicom_report_text(path), "dicom-report"
    except Exception:
        return _read_text_file(path), "text"


def _sentence_for_match(text: str, start: int, end: int) -> str:
    left_candidates = [text.rfind(mark, 0, start) for mark in ".!?\n"]
    right_candidates = [text.find(mark, end) for mark in ".!?\n"]
    left = max(left_candidates) + 1
    valid_right = [value for value in right_candidates if value != -1]
    right = min(valid_right) + 1 if valid_right else len(text)
    sentence = re.sub(r"\s+", " ", text[left:right]).strip()
    return sentence[:260]


def _is_negated_finding(text: str, start: int) -> bool:
    prefix = text[max(0, start - 42) : start].lower()
    return bool(
        re.search(
            r"\b(no|without|negative for|free of|absence of|absent|not seen|not identified|no evidence of)\s+"
            r"([a-z]+\s+){0,4}$",
            prefix,
        )
    )


def _detect_report_findings(text: str) -> List[dict]:
    findings = []
    seen = set()

    for severity, term in IMPORTANT_FINDING_RULES:
        pattern = re.compile(rf"\b{re.escape(term)}\w*\b", re.IGNORECASE)
        matches = [match for match in pattern.finditer(text) if not _is_negated_finding(text, match.start())]
        if not matches:
            continue
        first = matches[0]
        key = term.lower()
        if key in seen:
            continue
        seen.add(key)
        findings.append(
            {
                "term": term,
                "severity": severity,
                "count": len(matches),
                "context": _sentence_for_match(text, first.start(), first.end()),
            }
        )

    severity_order = {"critical": 0, "high": 1, "moderate": 2}
    findings.sort(key=lambda item: (severity_order.get(item["severity"], 9), item["term"]))
    return findings[:24]


def _nifti_to_zyx(path: Path) -> Tuple[np.ndarray, Tuple[float, float, float]]:
    image = nib.load(str(path))
    data = np.asarray(image.get_fdata(dtype=np.float32))

    if data.ndim == 4:
        data = data[..., 0]
    if data.ndim != 3:
        raise ValueError(f"NIfTI volume must be 3D or 4D, got shape {data.shape}.")

    zooms = image.header.get_zooms()[:3]
    volume = np.transpose(data, (2, 1, 0)).astype(np.float32, copy=False)
    spacing = (float(zooms[2]), float(zooms[1]), float(zooms[0]))
    return np.ascontiguousarray(volume), spacing


def _nifti_mask_to_zyx(path: Path) -> np.ndarray:
    image = nib.load(str(path))
    raw = np.asarray(image.dataobj)
    if raw.ndim == 4:
        raw = raw[..., 0]
    if raw.ndim != 3:
        raise ValueError(f"Segmentation mask must be 3D or 4D, got shape {raw.shape}.")

    mask = np.rint(np.transpose(raw, (2, 1, 0))).astype(np.uint16, copy=False)
    return np.ascontiguousarray(mask)


def _dicom_sort_key(dataset: pydicom.Dataset) -> float:
    try:
        position = dataset.ImagePositionPatient
        orientation = dataset.ImageOrientationPatient
        row = np.array([float(v) for v in orientation[:3]], dtype=np.float64)
        col = np.array([float(v) for v in orientation[3:]], dtype=np.float64)
        normal = np.cross(row, col)
        return float(np.dot(normal, np.array([float(v) for v in position], dtype=np.float64)))
    except Exception:
        return float(getattr(dataset, "InstanceNumber", 0))


def _find_dicom_series(folder: Path) -> List[Path]:
    series: Dict[str, List[Path]] = {}

    for candidate in folder.rglob("*"):
        if not candidate.is_file():
            continue
        try:
            ds = pydicom.dcmread(str(candidate), stop_before_pixels=True, force=False)
        except (InvalidDicomError, PermissionError, OSError):
            continue

        series_id = str(getattr(ds, "SeriesInstanceUID", "unknown"))
        series.setdefault(series_id, []).append(candidate)

    if not series:
        raise ValueError(f"No readable DICOM files found in {folder}.")

    return max(series.values(), key=len)


def _load_dicom_folder(folder: Path) -> Tuple[np.ndarray, Tuple[float, float, float]]:
    files = _find_dicom_series(folder)
    slices = []

    for path in files:
        ds = pydicom.dcmread(str(path), force=False)
        pixels = ds.pixel_array.astype(np.float32)
        slope = float(getattr(ds, "RescaleSlope", 1.0))
        intercept = float(getattr(ds, "RescaleIntercept", 0.0))
        pixels = pixels * slope + intercept
        slices.append((ds, pixels))

    slices.sort(key=lambda item: _dicom_sort_key(item[0]))

    volume = np.stack([pixels for _, pixels in slices], axis=0).astype(np.float32, copy=False)
    first = slices[0][0]
    pixel_spacing = getattr(first, "PixelSpacing", [1.0, 1.0])
    y_spacing = float(pixel_spacing[0])
    x_spacing = float(pixel_spacing[1])

    if len(slices) > 1:
        sorted_positions = [_dicom_sort_key(ds) for ds, _ in slices]
        diffs = np.diff(sorted_positions)
        z_spacing = float(np.median(np.abs(diffs))) if np.any(diffs) else float(getattr(first, "SliceThickness", 1.0))
    else:
        z_spacing = float(getattr(first, "SliceThickness", 1.0))

    return np.ascontiguousarray(volume), (z_spacing, y_spacing, x_spacing)


def _load_volume(path: Path) -> Tuple[np.ndarray, Tuple[float, float, float], str]:
    if path.is_dir():
        volume, spacing = _load_dicom_folder(path)
        return volume, spacing, "dicom"

    lower_name = path.name.lower()
    if lower_name.endswith(".nii") or lower_name.endswith(".nii.gz"):
        volume, spacing = _nifti_to_zyx(path)
        return volume, spacing, "nifti"

    raise ValueError("Unsupported volume path. Use a DICOM folder, .nii, or .nii.gz file.")


def _robust_min_max(data: np.ndarray) -> Tuple[float, float]:
    finite = data[np.isfinite(data)]
    if finite.size == 0:
        return 0.0, 1.0
    lo, hi = np.percentile(finite, [1, 99])
    if float(lo) == float(hi):
        lo, hi = float(np.min(finite)), float(np.max(finite))
    if float(lo) == float(hi):
        hi = float(lo) + 1.0
    return float(lo), float(hi)


def _class_color(label: int) -> str:
    palette = [
        "#00e5ff",
        "#ff4ecd",
        "#ffe66d",
        "#59ff9f",
        "#ff7a59",
        "#9d7cff",
        "#6df7c1",
        "#ff3366",
        "#f7f25c",
        "#55aaff",
    ]
    return palette[(int(label) - 1) % len(palette)]


def _serialize_volume(volume: Volume) -> dict:
    z, y, x = volume.data.shape
    return {
        "id": volume.volume_id,
        "source": volume.source,
        "kind": volume.kind,
        "shape": [int(z), int(y), int(x)],
        "spacing": [float(v) for v in volume.spacing],
        "dtype": "float32",
        "intensityMin": volume.intensity_min,
        "intensityMax": volume.intensity_max,
        "segmentations": [_serialize_segmentation(seg) for seg in volume.segmentations.values()],
    }


def _serialize_segmentation(segmentation: Segmentation) -> dict:
    return {
        "id": segmentation.mask_id,
        "source": segmentation.source,
        "shape": [int(v) for v in segmentation.data.shape],
        "dtype": "uint16",
        "classes": segmentation.classes,
    }


@app.route("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.route("/api/health")
def health():
    return jsonify({"ok": True})


@app.route("/api/load-volume", methods=["POST"])
def load_volume():
    try:
        path = _resolve_user_path(request.json.get("path", ""))
        data, spacing, kind = _load_volume(path)
        intensity_min, intensity_max = _robust_min_max(data)

        volume_id = uuid.uuid4().hex
        volume = Volume(
            volume_id=volume_id,
            source=str(path),
            kind=kind,
            data=data,
            spacing=spacing,
            intensity_min=intensity_min,
            intensity_max=intensity_max,
        )
        VOLUMES[volume_id] = volume
        return jsonify(_serialize_volume(volume))
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400


@app.route("/api/load-report", methods=["POST"])
def load_report():
    try:
        path = _resolve_user_path(request.json.get("path", ""))
        if not path.is_file():
            raise ValueError("Report path must be a file.")
        text, report_type = _extract_report_text(path)
        text = re.sub(r"\r\n?", "\n", text).strip()
        if not text:
            raise ValueError("No readable text found in this report.")

        return jsonify(
            {
                "source": str(path),
                "type": report_type,
                "text": text,
                "findings": _detect_report_findings(text),
            }
        )
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400


@app.route("/api/volume/<volume_id>/data")
def volume_data(volume_id: str):
    volume = VOLUMES.get(volume_id)
    if volume is None:
        return jsonify({"error": "Volume not found."}), 404

    payload = np.ascontiguousarray(volume.data.astype(np.float32, copy=False)).tobytes()
    return Response(payload, mimetype="application/octet-stream")


@app.route("/api/volume/<volume_id>/segmentation", methods=["POST"])
def load_segmentation(volume_id: str):
    volume = VOLUMES.get(volume_id)
    if volume is None:
        return jsonify({"error": "Load a volume before loading a segmentation."}), 404

    try:
        path = _resolve_user_path(request.json.get("path", ""))
        lower_name = path.name.lower()
        if not (lower_name.endswith(".nii") or lower_name.endswith(".nii.gz")):
            raise ValueError("Segmentation mask must be a .nii or .nii.gz file.")

        mask = _nifti_mask_to_zyx(path)
        if mask.shape != volume.data.shape:
            raise ValueError(f"Mask shape {mask.shape} does not match volume shape {volume.data.shape}.")

        labels = [int(v) for v in np.unique(mask) if int(v) != 0]
        classes = [
            {
                "label": label,
                "name": f"Class {label}",
                "color": _class_color(label),
                "opacity": 0.55,
            }
            for label in labels
        ]

        mask_id = uuid.uuid4().hex
        segmentation = Segmentation(mask_id=mask_id, source=str(path), data=mask, classes=classes)
        volume.segmentations[mask_id] = segmentation
        return jsonify(_serialize_segmentation(segmentation))
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400


@app.route("/api/volume/<volume_id>/segmentation/<mask_id>/data")
def segmentation_data(volume_id: str, mask_id: str):
    volume = VOLUMES.get(volume_id)
    if volume is None:
        return jsonify({"error": "Volume not found."}), 404
    segmentation = volume.segmentations.get(mask_id)
    if segmentation is None:
        return jsonify({"error": "Segmentation not found."}), 404

    payload = np.ascontiguousarray(segmentation.data.astype(np.uint16, copy=False)).tobytes()
    return Response(payload, mimetype="application/octet-stream")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5055"))
    host = os.environ.get("HOST", "127.0.0.1")
    app.run(host=host, port=port, debug=True)
